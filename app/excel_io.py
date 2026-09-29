"""Excel import/export. The data workbook doubles as the import template."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .checker import Snapshot
from .db import DAY_SHORT, DAYS, LESSON_KINDS, ROOM_KINDS, WEEKS, Database
YES = {"да", "yes", "1", "true", "+", "x", "✓"}
HEADER_FILL = PatternFill("solid", fgColor="DCE6F5")

SHEETS = {
    "Оборудование": ["Код", "Название"],
    "Аудитории": ["Номер", "Корпус", "Этаж", "Мест", "Тип", "ПК", "Оборудование", "Доступная среда", "Активна", "Примечание"],
    "Закрытия": ["Аудитория", "День", "Пара", "Причина"],
    "Преподаватели": ["ФИО", "Кафедра", "Доступная среда", "Недоступен", "Пожелания"],
    "Группы": ["Название", "Курс", "Студентов", "Корпус"],
    "Занятия": ["Дисциплина", "Тип", "Преподаватель", "Группы", "День", "Пара", "Неделя",
                "Оборудование", "Мин. ПК", "Аудитория", "Закреплено", "Время закреплено"],
}

PREF_HELP = ("Пожелания через «;»: корпус=Б:8 · избегать корпус=В:5 · аудитория=А-101:3 · "
             "оборудование=Проектор:4 · этаж<=2:5 · день=Пн:5 · избегать пара=1:6 "
             "(число после «:» — важность 1–10)")


OPTIONAL_COLUMNS = ("Примечание", "Время закреплено")


class ImportError_(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors[:5]))
        self.errors = errors


# ---------- helpers ----------
def _yn(v) -> bool:
    return str(v or "").strip().lower() in YES


def _bool_text(b) -> str:
    return "да" if b else ""


def _reverse(d: dict) -> dict:
    return {v.lower(): k for k, v in d.items()} | {k.lower(): k for k in d}


def _split(v) -> list[str]:
    return [s.strip() for s in re.split(r"[,;\n]", str(v or "")) if s.strip()]


def _day(v) -> int | None:
    s = str(v or "").strip().lower()
    if not s:
        return None
    for i, (short, full) in enumerate(zip(DAY_SHORT, DAYS)):
        if s in (short.lower(), full.lower(), str(i + 1)):
            return i
    raise ValueError(f"неизвестный день «{v}»")


def _int(v, default=0) -> int:
    if v is None or str(v).strip() == "":
        return default
    return int(float(str(v).replace(",", ".")))


def _slots_text(slots) -> str:
    return ", ".join(f"{DAY_SHORT[d]}{p}" for d, p in slots)


def _parse_slots(v) -> list[list[int]]:
    out = []
    for part in _split(v):
        m = re.fullmatch(r"([А-Яа-яA-Za-z]+)\s*(\d+)", part)
        if not m:
            raise ValueError(f"не понял слот «{part}» (пример: Сб1)")
        out.append([_day(m.group(1)), int(m.group(2))])
    return out


def prefs_to_text(prefs: list[dict], eq_names: dict) -> str:
    parts = []
    for p in prefs:
        w = p.get("weight", 5)
        if p["kind"] == "max_floor":
            parts.append(f"этаж<={p['value']}:{w}")
            continue
        avoid = "избегать " if p.get("mode") == "avoid" else ""
        if p["kind"] == "day":
            parts.append(f"{avoid}день={DAY_SHORT[int(p['value'])]}:{w}")
            continue
        if p["kind"] == "pair":
            parts.append(f"{avoid}пара={p['value']}:{w}")
            continue
        head = {"building": "корпус", "room": "аудитория", "equipment": "оборудование"}[p["kind"]]
        value = eq_names.get(p["value"], p["value"]) if p["kind"] == "equipment" else p["value"]
        if p["kind"] == "equipment" and p["value"] == "computers":
            value = "компьютеры"
        parts.append(f"{'избегать ' if p.get('mode') == 'avoid' else ''}{head}={value}:{w}")
    return "; ".join(parts)


def text_to_prefs(text, eq_by_name: dict) -> list[dict]:
    prefs = []
    for part in [s.strip() for s in str(text or "").split(";") if s.strip()]:
        m = re.fullmatch(r"(?i)(избегать\s+)?(корпус|аудитория|оборудование|этаж|день|пара)\s*(<=|=)\s*([^:]+?)\s*(?::\s*(\d+))?", part)
        if not m:
            raise ValueError(f"не понял пожелание «{part}»")
        avoid, head, _, value, weight = m.groups()
        weight = max(1, min(10, int(weight or 5)))
        head = head.lower()
        if head == "этаж":
            prefs.append({"kind": "max_floor", "value": int(value), "mode": "prefer", "weight": weight})
            continue
        if head == "день":
            prefs.append({"kind": "day", "value": _day(value), "mode": "avoid" if avoid else "prefer", "weight": weight})
            continue
        if head == "пара":
            prefs.append({"kind": "pair", "value": int(value), "mode": "avoid" if avoid else "prefer", "weight": weight})
            continue
        kind = {"корпус": "building", "аудитория": "room", "оборудование": "equipment"}[head]
        if kind == "equipment":
            value = "computers" if value.lower() in ("компьютеры", "пк") else eq_by_name.get(value.lower(), value)
        prefs.append({"kind": kind, "value": value, "mode": "avoid" if avoid else "prefer", "weight": weight})
    return prefs


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:300]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width))
    ws.freeze_panes = "A2"
    return ws


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- export data ----------
def export_data(db: Database) -> bytes:
    s = Snapshot(db)
    eq = s.equipment_names
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Оборудование", SHEETS["Оборудование"], [[k, v] for k, v in eq.items()])
    _sheet(wb, "Аудитории", SHEETS["Аудитории"], [
        [r["name"], r["building"], r["floor"], r["capacity"], ROOM_KINDS.get(r["kind"], r["kind"]), r["computers"],
         ", ".join(eq.get(e, e) for e in r["equipment"]), _bool_text(r["accessible"]), _bool_text(r["active"]), r["notes"]]
        for r in s.rooms.values()])
    blocks = []
    for room_id, items in s.blocks.items():
        for b in items:
            blocks.append([s.rooms[room_id]["name"], DAY_SHORT[b["day"]] if b["day"] is not None else "",
                           b["pair"] or "", b["reason"]])
    _sheet(wb, "Закрытия", SHEETS["Закрытия"], blocks)
    ws = _sheet(wb, "Преподаватели", SHEETS["Преподаватели"], [
        [t["name"], t["department"], _bool_text(t["needs_accessible"]), _slots_text(t["unavailable"]),
         prefs_to_text(t["preferences"], eq)] for t in s.teachers.values()])
    ws.cell(row=1, column=7, value=PREF_HELP).font = Font(italic=True, color="666666")
    _sheet(wb, "Группы", SHEETS["Группы"], [
        [g["name"], g["course"], g["size"], g["home_building"]] for g in s.groups.values()])
    _sheet(wb, "Занятия", SHEETS["Занятия"], [
        [l["subject"], LESSON_KINDS.get(l["kind"], l["kind"]),
         s.teachers[l["teacher_id"]]["name"] if l["teacher_id"] in s.teachers else "",
         ", ".join(s.groups[g]["name"] for g in l["group_ids"] if g in s.groups),
         DAY_SHORT[l["day"]], l["pair"], WEEKS.get(l["week"], l["week"]),
         ", ".join(eq.get(e, e) for e in l["required_equipment"]), l["min_computers"] or "",
         s.rooms[l["room_id"]]["name"] if l["room_id"] in s.rooms else "", _bool_text(l["locked"]),
         _bool_text(l["fixed_time"])]
        for l in s.lessons.values()])
    return _save(wb)


# ---------- import data ----------
def import_data(db: Database, content: bytes) -> dict:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    errors: list[str] = []

    def rows(title):
        if title not in wb.sheetnames:
            return []
        ws = wb[title]
        header = [str(c.value or "").strip() for c in ws[1]]
        idx = {h: i for i, h in enumerate(header)}
        missing = [h for h in SHEETS[title] if h not in idx and h not in OPTIONAL_COLUMNS]
        if missing:
            errors.append(f"лист «{title}»: нет столбцов {', '.join(missing)}")
            return []
        out = []
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(v not in (None, "") for v in row):
                continue
            out.append((n, {h: (row[i] if i < len(row) else None) for h, i in idx.items()}))
        return out

    def guard(title, n, fn):
        try:
            return fn()
        except (ValueError, KeyError, TypeError) as e:
            errors.append(f"лист «{title}», строка {n}: {e}")
            return None

    # equipment
    equipment = {e["code"]: e["name"] for e in db.equipment()}
    for n, r in rows("Оборудование"):
        code, name = str(r["Код"] or "").strip(), str(r["Название"] or "").strip()
        if code:
            equipment[code] = name or code
    eq_by_name = {v.lower(): k for k, v in equipment.items()} | {k.lower(): k for k in equipment}

    def eq_codes(v):
        out = []
        for name in _split(v):
            if name.lower() not in eq_by_name:
                raise ValueError(f"неизвестное оборудование «{name}» (добавьте на лист «Оборудование»)")
            out.append(eq_by_name[name.lower()])
        return out

    room_kind = _reverse(ROOM_KINDS)
    lesson_kind = _reverse(LESSON_KINDS)
    week_kind = _reverse(WEEKS) | {"каждую": "all", "все": "all", "": "all", "числ": "odd", "знам": "even"}

    rooms = {}
    for n, r in rows("Аудитории"):
        def parse(r=r):
            name = str(r["Номер"]).strip()
            kind = room_kind.get(str(r["Тип"] or "практическая").strip().lower())
            if not kind:
                raise ValueError(f"тип аудитории «{r['Тип']}»: {', '.join(ROOM_KINDS.values())}")
            return name, dict(name=name, building=str(r["Корпус"] or "").strip(), floor=_int(r["Этаж"], 1),
                              capacity=_int(r["Мест"]), kind=kind, computers=_int(r["ПК"]),
                              equipment=eq_codes(r["Оборудование"]), accessible=_yn(r["Доступная среда"]),
                              active=True if r["Активна"] in (None, "") else _yn(r["Активна"]),
                              notes=str(r.get("Примечание") or ""))
        res = guard("Аудитории", n, parse)
        if res:
            if res[0] in rooms:
                errors.append(f"лист «Аудитории», строка {n}: повтор номера {res[0]}")
            rooms[res[0]] = res[1]

    blocks = []
    for n, r in rows("Закрытия"):
        def parse(r=r):
            name = str(r["Аудитория"]).strip()
            if name not in rooms:
                raise ValueError(f"нет аудитории «{name}»")
            pair = _int(r["Пара"], 0) or None
            return dict(room=name, day=_day(r["День"]), pair=pair, reason=str(r["Причина"] or ""))
        res = guard("Закрытия", n, parse)
        if res:
            blocks.append(res)

    teachers = {}
    for n, r in rows("Преподаватели"):
        def parse(r=r):
            name = str(r["ФИО"]).strip()
            return name, dict(name=name, department=str(r["Кафедра"] or ""),
                              needs_accessible=_yn(r["Доступная среда"]),
                              unavailable=_parse_slots(r["Недоступен"]),
                              preferences=text_to_prefs(r["Пожелания"], eq_by_name))
        res = guard("Преподаватели", n, parse)
        if res:
            teachers[res[0]] = res[1]

    groups = {}
    for n, r in rows("Группы"):
        def parse(r=r):
            name = str(r["Название"]).strip()
            return name, dict(name=name, course=_int(r["Курс"], 1), size=_int(r["Студентов"]),
                              home_building=str(r["Корпус"] or "").strip())
        res = guard("Группы", n, parse)
        if res:
            if res[0] in groups:
                errors.append(f"лист «Группы», строка {n}: повтор группы {res[0]}")
            groups[res[0]] = res[1]

    lessons = []
    for n, r in rows("Занятия"):
        def parse(r=r):
            kind = lesson_kind.get(str(r["Тип"] or "практика").strip().lower())
            if not kind:
                raise ValueError(f"тип занятия «{r['Тип']}»: {', '.join(LESSON_KINDS.values())}")
            week = week_kind.get(str(r["Неделя"] or "").strip().lower())
            if not week:
                raise ValueError(f"неделя «{r['Неделя']}»: каждую / числитель / знаменатель")
            teacher = str(r["Преподаватель"] or "").strip()
            if teacher and teacher not in teachers:
                raise ValueError(f"нет преподавателя «{teacher}»")
            gnames = _split(r["Группы"])
            unknown = [g for g in gnames if g not in groups]
            if unknown:
                raise ValueError(f"нет групп: {', '.join(unknown)}")
            room = str(r["Аудитория"] or "").strip()
            if room and room not in rooms:
                raise ValueError(f"нет аудитории «{room}»")
            day = _day(r["День"])
            if day is None:
                raise ValueError("не указан день")
            return dict(subject=str(r["Дисциплина"]).strip(), kind=kind, teacher=teacher or None, groups=gnames,
                        day=day, pair=_int(r["Пара"], 1), week=week, required_equipment=eq_codes(r["Оборудование"]),
                        min_computers=_int(r["Мин. ПК"]), room=room or None, locked=_yn(r["Закреплено"]),
                        fixed_time=_yn(r.get("Время закреплено")))
        res = guard("Занятия", n, parse)
        if res:
            lessons.append(res)

    if errors:
        raise ImportError_(errors)

    # write everything
    db.clear()
    for code, name in equipment.items():
        db.upsert_equipment(code, name)
    room_ids = {name: db.insert("rooms", d) for name, d in rooms.items()}
    for b in blocks:
        db.insert("room_blocks", dict(room_id=room_ids[b["room"]], day=b["day"], pair=b["pair"], reason=b["reason"]))
    teacher_ids = {name: db.insert("teachers", d) for name, d in teachers.items()}
    group_ids = {name: db.insert("student_groups", d) for name, d in groups.items()}
    for l in lessons:
        db.save_lesson(dict(subject=l["subject"], kind=l["kind"], teacher_id=teacher_ids.get(l["teacher"]),
                            day=l["day"], pair=l["pair"], week=l["week"], required_equipment=l["required_equipment"],
                            min_computers=l["min_computers"], room_id=room_ids.get(l["room"]), locked=l["locked"],
                            fixed_time=l["fixed_time"],
                            group_ids=[group_ids[g] for g in l["groups"]]))
    return {"rooms": len(rooms), "teachers": len(teachers), "groups": len(groups), "lessons": len(lessons)}


# ---------- export schedule ----------
def export_schedule(db: Database) -> bytes:
    s = Snapshot(db)
    times = {p["n"]: f"{p['start']}–{p['end']}" for p in s.settings["pairs"]}
    wb = Workbook()
    wb.remove(wb.active)
    lessons = sorted(s.lessons.values(), key=lambda l: (l["day"], l["pair"], l["subject"]))

    def room_name(l):
        return s.rooms[l["room_id"]]["name"] if l["room_id"] in s.rooms else "—"

    def teacher_name(l):
        return s.teachers[l["teacher_id"]]["name"] if l["teacher_id"] in s.teachers else ""

    def week_mark(l):
        return {"odd": " [числ.]", "even": " [знам.]"}.get(l["week"], "")

    _sheet(wb, "Список", ["День", "Пара", "Время", "Неделя", "Дисциплина", "Тип", "Преподаватель", "Группы",
                          "Аудитория", "Корпус"], [
        [DAYS[l["day"]], l["pair"], times.get(l["pair"], ""), WEEKS[l["week"]], l["subject"],
         LESSON_KINDS.get(l["kind"], l["kind"]), teacher_name(l),
         ", ".join(s.groups[g]["name"] for g in l["group_ids"] if g in s.groups), room_name(l),
         s.rooms[l["room_id"]]["building"] if l["room_id"] in s.rooms else ""] for l in lessons])

    def grid(title, columns, key_of, text_of):
        ws = wb.create_sheet(title)
        ws.append(["День", "Пара"] + [c[1] for c in columns])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
        cells: dict[tuple, list[str]] = {}
        for l in lessons:
            for k in key_of(l):
                cells.setdefault((l["day"], l["pair"], k), []).append(text_of(l) + week_mark(l))
        for d in range(s.settings.get("days", 6)):
            for p in s.pairs:
                ws.append([DAY_SHORT[d], f"{p} ({times.get(p, '')})"] +
                          ["\n".join(cells.get((d, p, c[0]), [])) for c in columns])
        for i in range(3, len(columns) + 3):
            ws.column_dimensions[get_column_letter(i)].width = 28
        ws.column_dimensions["B"].width = 16
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"

    grid("По группам", [(g["id"], g["name"]) for g in s.groups.values()], lambda l: l["group_ids"],
         lambda l: f"{l['subject']} ({LESSON_KINDS.get(l['kind'], '')[:3].lower()}.) {room_name(l)}, {teacher_name(l)}")
    grid("По аудиториям", [(r["id"], r["name"]) for r in s.rooms.values()], lambda l: [l["room_id"]],
         lambda l: f"{l['subject']}: " + ", ".join(s.groups[g]["name"] for g in l["group_ids"] if g in s.groups))
    grid("По преподавателям", [(t["id"], t["name"]) for t in s.teachers.values()], lambda l: [l["teacher_id"]],
         lambda l: f"{l['subject']}, {room_name(l)}: " + ", ".join(s.groups[g]["name"] for g in l["group_ids"] if g in s.groups))
    return _save(wb)
